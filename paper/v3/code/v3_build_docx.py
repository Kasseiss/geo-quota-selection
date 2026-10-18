"""Write the Word submission on the official MDPI Sensors Word template (assets/mdpi_sensors_template.docx, converted
unchanged from template/sensors-template.dot).

The template's own front page (article type, title, author block, editor/dates table, affiliations, Highlights, Abstract,
Keywords) is kept and filled in place; its instructional body is removed and the article is appended with the MDPI
paragraph styles. Author names and affiliations stay as the template placeholders for the authors to complete.
"""
import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import v3_text as T
from v3_common import ROOT, OUT
from v3_content import REFS, ABBREVIATIONS, Citer, outline
from v3_tables import _border, _margins

TEMPLATE = Path(__file__).resolve().parent / 'assets/mdpi_sensors_template.docx'
DOCX = ROOT / 'paper/v3/submission/manuscript.docx'
FULL = 7.2   # usable page width (inches) for full-width figures and tables


def set_text(p, text, keep_first_runs=1):
    runs = p.runs
    for r in runs[keep_first_runs:]:
        r._r.getparent().remove(r._r)
    runs[keep_first_runs - 1].text = text


def add_p(doc, style, text='', citer=None):
    p = doc.add_paragraph(style=style)
    if text:
        p.add_run(citer.numbers(text) if citer else text)
    return p


def figure(doc, name, caption, n, citer):
    p = doc.add_paragraph(style='MDPI_5.2_figure')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(OUT / f'{name}.png'), width=Inches(6.9))
    c = doc.add_paragraph(style='MDPI_5.1_figure_caption')
    a = c.add_run(f'Figure {n}. '); a.bold = True
    c.add_run(citer.numbers(caption))


def table(doc, spec, n):
    cap = doc.add_paragraph(style='MDPI_4.1_table_caption')
    a = cap.add_run(f'Table {n}. '); a.bold = True
    cap.add_run(spec['caption'])
    cap.paragraph_format.keep_with_next = True
    hdrs, rows, widths = spec['headers'], spec['rows'], spec['widths']
    scale = min(1.0, FULL / sum(widths)) if sum(widths) > FULL else FULL / sum(widths) if sum(widths) > 5.4 else 1.0
    widths = [w * scale for w in widths]
    ncol = len(hdrs[-1])
    t = doc.add_table(rows=0, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    def write(cell, text, align, bold=False):
        p = cell.paragraphs[0]
        p.style = doc.styles['MDPI_4.2_table_body']
        p.alignment = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'c': WD_ALIGN_PARAGRAPH.CENTER}[align]
        r = p.add_run(str(text)); r.bold = bold

    def finish(row, top=None, bottom=None):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for k, c in enumerate(row.cells):
            c.width = Inches(widths[k]); _margins(c); _border(c, top=top, bottom=bottom)

    for h, hdr in enumerate(hdrs):
        row = t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        start = 0
        for k in range(ncol):
            if hdr[k] is not None:
                start = k
                write(row.cells[k], hdr[k], 'l' if k == 0 else 'c', bold=True)
            else:
                row.cells[start].merge(row.cells[k])
        last = h == len(hdrs) - 1
        finish(row, top=12 if h == 0 else None, bottom=6 if last else None)
        if not last:
            for k in range(1, ncol):
                if hdr[k] is not None:
                    _border(row.cells[k], top=12, bottom=4)
    body = []
    for i, r in enumerate(rows):
        row = t.add_row()
        for k, v in enumerate(r):
            write(row.cells[k], v, spec['align'][k], bold=(i, k) in spec['bold'])
        body.append(row)
    for i, row in enumerate(body):
        finish(row, bottom=12 if i == len(body) - 1 else None)
    foot = doc.add_paragraph(style='MDPI_4.3_table_footer')
    foot.add_run(spec['notes'])


def reference(doc, runs, template_ppr):
    p = doc.add_paragraph(style='MDPI_8.1_references')
    old = p._p.find(qn('w:pPr'))
    if old is not None:
        p._p.remove(old)
    p._p.insert(0, copy.deepcopy(template_ppr))
    for text, style in runs:
        r = p.add_run(text)
        r.italic = style == 'i'
        r.bold = style == 'b'


def main():
    doc = Document(TEMPLATE)
    body = doc.element.body
    els = list(body.iterchildren())
    ref_ppr = copy.deepcopy(next(e for e in els if 'Author 1, A.B.; Author 2, C.D. Title of the article'
                                 in ''.join(e.itertext())).find(qn('w:pPr')))
    paras = doc.paragraphs

    # ---- front page, filled in place
    set_text(paras[0], 'Article')
    set_text(paras[1], T.TITLE)
    hl = [p for p in paras[:18]]
    expl = next(p for p in hl if p.text.startswith('This is an optional section'))
    expl._p.getparent().remove(expl._p)
    bullets = [p for p in doc.paragraphs[:18] if p.text in ('First bullet.', 'Second bullet.')]
    for p, text in zip(bullets, T.HIGHLIGHTS_FINDINGS + T.HIGHLIGHTS_IMPLICATIONS):
        set_text(p, text)
    abstract = next(p for p in doc.paragraphs[:18] if p.text.startswith('A single paragraph'))
    set_text(abstract, T.ABSTRACT)
    kw = next(p for p in doc.paragraphs[:18] if p.text.startswith('Keywords'))
    for r in kw.runs[1:]:
        r._r.getparent().remove(r._r)
    kw.runs[0].text = 'Keywords: '
    kw.runs[0].bold = True
    kw.add_run(T.KEYWORDS)

    # ---- remove the instructional body (everything after the line under the keywords)
    els = list(body.iterchildren())
    line_p = next(p for p in doc.paragraphs if p.style.name == 'MDPI_1.9_line')._p
    line = els.index(line_p)
    for e in els[line + 1:]:
        if e.tag != qn('w:sectPr'):
            body.remove(e)

    # ---- article body
    citer = Citer()
    sec = sub = fig_n = tab_n = 0
    for b in outline():
        kind = b[0]
        if kind == 'h1':
            sec += 1; sub = 0
            add_p(doc, 'MDPI_2.1_heading1', f'{sec}. {b[1]}')
        elif kind == 'h2':
            sub += 1
            add_p(doc, 'MDPI_2.2_heading2', f'{sec}.{sub}. {b[1]}')
        elif kind == 'p':
            add_p(doc, 'MDPI_3.1_text', b[1], citer)
        elif kind == 'enum':
            for x in b[1]:
                add_p(doc, 'MDPI_3.7_itemize', x, citer)
        elif kind == 'fig':
            fig_n += 1
            figure(doc, b[1], b[2], fig_n, citer)
        elif kind == 'table':
            tab_n += 1
            table(doc, b[1], tab_n)

    # ---- back matter
    back = dict(T.BACK)
    items = [('Supplementary Materials', T.SUPP), ('Author Contributions', back['Author Contributions']),
             ('Funding', back['Funding']), ('Institutional Review Board Statement', 'Not applicable.'),
             ('Informed Consent Statement', 'Not applicable.'),
             ('Data Availability Statement', back['Data Availability Statement']),
             ('Acknowledgments', back['Use of Generative AI']),
             ('Conflicts of Interest', back['Conflicts of Interest'])]
    for h, text in items:
        p = doc.add_paragraph(style='MDPI_6.2_back_matter')
        a = p.add_run(h + ': '); a.bold = True
        p.add_run(text)
    add_p(doc, 'MDPI_2.1_heading1', 'Abbreviations')
    add_p(doc, 'MDPI_3.2_text_no_indent', 'The following abbreviations are used in this manuscript:')
    ab = doc.add_table(rows=0, cols=2)
    for a, b in ABBREVIATIONS:
        row = ab.add_row()
        for k, v in enumerate((a, b)):
            row.cells[k].paragraphs[0].style = doc.styles['MDPI_4.2_table_body']
            row.cells[k].paragraphs[0].add_run(v)
            row.cells[k].width = Inches(1.0 if k == 0 else 3.5)
    add_p(doc, 'MDPI_2.1_heading1', 'References')
    for key in citer.order:
        reference(doc, REFS[key], ref_ppr)
    unused = set(REFS) - set(citer.order)
    assert not unused, unused
    DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX)
    print('saved', DOCX, '| figures', fig_n, '| tables', tab_n, '| references', len(citer.order))


if __name__ == '__main__':
    main()
