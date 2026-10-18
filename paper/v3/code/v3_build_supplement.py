"""Supplementary material for v3: the v2 supplement with the new article title, a naming key that links its rule labels
to the main text, and the fixed qualitative examples added as Figure S3. Nothing else in the v2 supplement changes."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

import v3_text as T
from v3_common import ROOT

SRC = ROOT / 'paper/v2/supplementary_material.docx'
DST = ROOT / 'paper/v3/submission/supplementary_materials.docx'
QUAL = ROOT / 'figures/v2/qualitative_fixed_examples.png'

NAMING = ('Naming used in this supplement and in the main text: K = visual k-center; GeoK = geo-quota k-center; '
          'Geo-random = geo-quota random; Permuted GeoK = placebo-quota k-center; weighted FL = geo-weighted facility '
          'location; permuted weighted FL = placebo-weighted facility location; FD adaptation = feature-cluster round '
          'robin; "active" = the 576-fit main matrix (90 complete-model GPU fits and 486 frozen-head fits); "pixel-budget '
          'control" = the equal-label experiment of the main text.')

CAPTION = ('Figure S3. Fixed qualitative examples. The first alphabetical image of the first two panel regions with a '
           'single unambiguous CC BY 4.0 catalogue source (Al Qurnah and Baybay) was chosen before these images or '
           'their per-image outcomes were inspected. Predictions come from the one archived complete-model checkpoint '
           '(CPU reference execution, source0, seeds (0, 0), geo-quota k-center); reloading it reproduced the stored '
           'confusion matrices of both images exactly. The examples illustrate one fit and are not a comparison between '
           'rules. Void pixels are light grey and are ignored in scoring. Imagery: OpenAerialMap / Open Imagery Network '
           'contributors (Bilal Ko\u00e7; SkyEye), CC BY 4.0, resampled to 448 \u00d7 448; masks from OpenEarthMap.')


def main():
    d = Document(SRC)
    title = d.paragraphs[0]
    for r in title.runs[1:]:
        r._r.getparent().remove(r._r)
    title.runs[0].text = 'Supplementary Materials for: ' + T.TITLE
    kw = d.paragraphs[1]
    p = d.paragraphs[2].insert_paragraph_before(NAMING, style=kw.style)
    s9 = next(p for p in d.paragraphs if p.text.startswith('S9. Data attribution'))
    pic = s9.insert_paragraph_before('')
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(QUAL), width=Inches(5.2))
    s9.insert_paragraph_before(CAPTION, style='Caption')
    DST.parent.mkdir(parents=True, exist_ok=True)
    d.save(DST)
    print('saved', DST)


if __name__ == '__main__':
    main()
