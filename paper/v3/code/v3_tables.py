"""Three-line (booktabs-style) tables for python-docx, in the layout used by MDPI journals."""
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ALIGN = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'c': WD_ALIGN_PARAGRAPH.CENTER, 'r': WD_ALIGN_PARAGRAPH.RIGHT}


def _border(cell, **edges):
    pr = cell._tc.get_or_add_tcPr()
    old = pr.find(qn('w:tcBorders'))
    if old is not None:
        pr.remove(old)
    b = OxmlElement('w:tcBorders'); pr.append(b)
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + edge)
        size = edges.get(edge)
        e.set(qn('w:val'), 'single' if size else 'nil')
        if size:
            e.set(qn('w:sz'), str(size)); e.set(qn('w:color'), '000000'); e.set(qn('w:space'), '0')
        b.append(e)


def _margins(cell, tb=40, lr=70):
    pr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for edge, v in (('top', tb), ('bottom', tb), ('left', lr), ('right', lr)):
        e = OxmlElement('w:' + edge); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    pr.append(m)


def _write(cell, text, align, size, bold=False, italic=False):
    p = cell.paragraphs[0]
    p.alignment = ALIGN[align]
    pf = p.paragraph_format
    pf.space_before = Pt(1); pf.space_after = Pt(1); pf.line_spacing = 1.0
    r = p.add_run(str(text))
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def caption(d, label, text):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    a = p.add_run(label + ' '); a.bold = True; a.font.size = Pt(9.5)
    b = p.add_run(text); b.font.size = Pt(9.5)
    return p


def booktabs(d, label, text, headers, rows, widths, align=None, notes=None, size=9, bold=None, groups=None):
    """rows may contain ('group', 'Label') tuples, rendered as italic spanning sub-headings.

    bold: set of (row_index, col_index) to set in bold, counted over data rows only.
    headers may be a list of rows (for a two-level header); a None entry merges with the cell to its left.
    """
    caption(d, label, text)
    header_rows = headers if isinstance(headers[0], (list, tuple)) else [headers]
    ncol = len(header_rows[-1])
    align = align or ['l'] + ['c'] * (ncol - 1)
    bold = bold or set()
    t = d.add_table(rows=0, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for k, w in enumerate(widths):
        t.columns[k].width = Inches(w)

    def finish(row, top=None, bottom=None):
        tr = row._tr.get_or_add_trPr()
        tr.append(OxmlElement('w:cantSplit'))
        for k, c in enumerate(row.cells):
            c.width = Inches(widths[k])
            _margins(c)
            _border(c, top=top, bottom=bottom)

    for h, hdr in enumerate(header_rows):
        row = t.add_row()
        el = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(el)
        start = 0
        for k in range(ncol):
            if hdr[k] is not None:
                start = k
                _write(row.cells[k], hdr[k], 'l' if k == 0 else 'c', size, bold=True)
            else:
                row.cells[start].merge(row.cells[k])
        last = h == len(header_rows) - 1
        finish(row, top=12 if h == 0 else None, bottom=6 if last else None)
        if not last:  # partial rule under spanning header cells
            for k in range(1, ncol):
                if hdr[k] is not None:
                    _border(row.cells[k], top=12 if h == 0 else None, bottom=4)
    if len(header_rows) > 1:  # merge first-column headers vertically when lower rows are blank
        hrows = t.rows[:len(header_rows)]
        for k in range(ncol):
            if all(hr[k] == '' for hr in header_rows[1:]):
                c = hrows[0].cells[k].merge(hrows[-1].cells[k])
                for extra in c.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _border(c, top=12, bottom=6)
    data_i = 0
    body = []
    for r in rows:
        row = t.add_row()
        if isinstance(r, tuple) and r and r[0] == 'group':
            cell = row.cells[0].merge(row.cells[ncol - 1])
            _write(cell, r[1], 'l', size, italic=True)
        else:
            for k, v in enumerate(r):
                _write(row.cells[k], v, align[k], size, bold=(data_i, k) in bold)
            data_i += 1
        body.append(row)
    for i, row in enumerate(body):
        finish(row, bottom=12 if i == len(body) - 1 else None)
        if i < len(body) - 1:
            for p in row.cells[0].paragraphs:
                p.paragraph_format.keep_with_next = len(rows) <= 8
    if notes:
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(notes); r.font.size = Pt(8)
    else:
        d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t
